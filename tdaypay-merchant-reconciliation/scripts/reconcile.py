"""Build a deterministic TDayPay merchant reconciliation JSON data set."""

from __future__ import annotations

import argparse
import csv
import glob
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo


ALIASES = {
    "平台订单号": "平台订单号", "platformOrderNo": "平台订单号", "platform_order_no": "平台订单号",
    "商户订单号": "商户订单号", "merchantOrderNo": "商户订单号", "merchant_order_no": "商户订单号",
    "订单金额": "订单金额", "amount": "订单金额", "orderAmount": "订单金额",
    "币种": "币种", "currency": "币种",
    "订单状态": "订单状态", "status": "订单状态", "orderStatus": "订单状态",
    "创建日期": "创建日期", "createDate": "创建日期", "createdDate": "创建日期",
    "创建时间": "创建时间", "createTime": "创建时间", "createdAt": "创建时间",
    "支付类型": "支付类型", "orderType": "支付类型", "paymentType": "支付类型",
    "结算状态": "结算状态", "settlementStatus": "结算状态",
    "流水ID": "流水ID", "flowId": "流水ID", "flow_id": "流水ID",
    "订单号": "订单号", "orderNo": "订单号", "orderId": "订单号",
    "资金域": "资金域", "fundingDomain": "资金域",
    "账户类型": "账户类型", "amountType": "账户类型", "accountType": "账户类型",
}

SYSTEM_REQUIRED = {"平台订单号", "商户订单号", "订单金额", "币种", "订单状态", "创建日期"}
TDAYPAY_REQUIRED = SYSTEM_REQUIRED | {"结算状态"}
FLOW_REQUIRED = {"流水ID", "订单号", "资金域", "账户类型", "创建时间"}
MONEY_FIELDS = {"订单金额", "出款手续费", "出款费率", "收款手续费", "收款费率", "期初余额", "发生额", "期末余额"}
IDENTIFIER_FIELDS = {"平台订单号", "商户订单号", "流水ID", "账户ID", "商户编号", "订单号", "银行单号", "银行关联号"}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).replace("\ufeff", "").strip()
    while len(text) >= 2 and text.startswith("'") and text.endswith("'"):
        text = text[1:-1].strip()
    if text.startswith("'"):
        text = text[1:]
    return text.strip()


def canonical_header(value: Any) -> str:
    header = clean(value)
    return ALIASES.get(header, header)


def normalize_date(value: Any) -> str:
    text = clean(value)
    digits = re.sub(r"[^0-9]", "", text)
    if len(digits) >= 8:
        candidate = digits[:8]
        try:
            datetime.strptime(candidate, "%Y%m%d")
            return candidate
        except ValueError:
            pass
    raise ValueError(f"无法识别日期值: {text!r}")


def normalize_money(value: Any) -> str:
    text = clean(value).replace(",", "")
    if text == "":
        return ""
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"无法识别金额值: {text!r}") from exc
    normalized = format(number, "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return normalized or "0"


def decimal_value(value: Any) -> Decimal:
    text = normalize_money(value)
    return Decimal(text or "0")


def normalize_order_type(value: Any) -> str:
    text = clean(value).upper()
    return {"PAYIN": "PAYMENT", "PAYMENT": "PAYMENT", "PAYOUT": "PAYOUT"}.get(text, text)


def normalize_row(row: dict[str, Any]) -> dict[str, str]:
    result: dict[str, str] = {}
    for raw_key, value in row.items():
        key = canonical_header(raw_key)
        if key in result and clean(value) and result[key] != clean(value):
            raise ValueError(f"列别名冲突: {raw_key!r} 映射到已有列 {key!r}")
        result[key] = clean(value)
    for key in MONEY_FIELDS & result.keys():
        result[key] = normalize_money(result[key])
    return result


def read_csv_file(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError(f"CSV 没有表头: {path}")
        return [normalize_row(dict(row)) for row in reader]


def read_excel_file(path: Path, sheet_name: str | None) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取 Excel 需要 openpyxl；请使用工作区依赖中的 Python。") from exc
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
    rows = sheet.iter_rows()
    try:
        header_cells = next(rows)
    except StopIteration as exc:
        raise ValueError(f"Excel 工作表为空: {path}#{sheet.title}") from exc
    headers = [canonical_header(cell.value) for cell in header_cells]
    if not any(headers):
        raise ValueError(f"Excel 工作表没有表头: {path}#{sheet.title}")
    result: list[dict[str, str]] = []
    for row_number, cells in enumerate(rows, 2):
        raw: dict[str, Any] = {}
        for header, cell in zip(headers, cells):
            if not header:
                continue
            if header in IDENTIFIER_FIELDS and cell.value not in (None, "") and not isinstance(cell.value, str):
                raise ValueError(
                    f"标识列必须是文本，避免精度丢失: {path}#{sheet.title}!{cell.coordinate} ({header})"
                )
            raw[header] = cell.value
        if any(clean(value) for value in raw.values()):
            result.append(normalize_row(raw))
    return result


def read_table(path: Path, sheet_name: str | None = None) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".txt"}:
        return read_csv_file(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_excel_file(path, sheet_name)
    raise ValueError(f"不支持的文件类型: {path}")


def expand_paths(values: Iterable[str]) -> list[Path]:
    paths: list[Path] = []
    for value in values:
        matches = [Path(p) for p in glob.glob(value)]
        if not matches:
            matches = [Path(value)]
        for path in matches:
            if path.is_dir():
                paths.extend(sorted(p for p in path.rglob("*") if p.suffix.lower() in {".csv", ".xlsx", ".xlsm"}))
            elif path.is_file():
                paths.append(path)
            else:
                raise FileNotFoundError(path)
    unique: list[Path] = []
    seen: set[Path] = set()
    for path in paths:
        resolved = path.resolve()
        if resolved not in seen:
            seen.add(resolved)
            unique.append(resolved)
    return unique


def require_columns(rows: list[dict[str, str]], required: set[str], label: str) -> list[str]:
    if not rows:
        raise ValueError(f"{label} 没有数据行")
    headers = list(rows[0].keys())
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"{label} 缺少必需列: {', '.join(missing)}")
    return headers


def load_many(paths: list[Path]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for path in paths:
        rows.extend(read_table(path))
    return rows


def filter_orders(
    rows: list[dict[str, str]], currency: str, start_key: str, end_key: str, label: str
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    included: list[dict[str, str]] = []
    excluded: list[dict[str, str]] = []
    for index, row in enumerate(rows, 2):
        reasons: list[str] = []
        row_currency = clean(row.get("币种")).upper()
        row_status = clean(row.get("订单状态")).upper()
        try:
            row_date = normalize_date(row.get("创建日期"))
        except ValueError:
            row_date = ""
            reasons.append("创建日期无效")
        if row_currency != currency:
            reasons.append(f"币种不是 {currency}")
        if row_status != "SUCCESS":
            reasons.append("订单状态不是 SUCCESS")
        if row_date and not (start_key <= row_date <= end_key):
            reasons.append("创建日期不在对账范围")
        if not clean(row.get("平台订单号")):
            reasons.append("平台订单号为空")
        if reasons:
            excluded.append({"来源": label, "源行号": str(index), "排除原因": "；".join(reasons), **row})
        else:
            normalized = dict(row)
            normalized["币种"] = row_currency
            normalized["订单状态"] = row_status
            normalized["创建日期"] = row_date
            if "支付类型" in normalized:
                normalized["支付类型"] = normalize_order_type(normalized["支付类型"])
            included.append(normalized)
    return included, excluded


def ensure_unique(rows: list[dict[str, str]], key: str, label: str) -> None:
    counts = Counter(clean(row.get(key)) for row in rows)
    duplicates = sorted(value for value, count in counts.items() if value and count > 1)
    if duplicates:
        sample = ", ".join(duplicates[:10])
        raise ValueError(f"{label} 的 {key} 重复: {sample}")


def deduplicate_flows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    by_id: dict[str, dict[str, str]] = {}
    for row in rows:
        flow_id = clean(row.get("流水ID"))
        if not flow_id:
            raise ValueError("记账流水存在空 流水ID")
        if flow_id in by_id and by_id[flow_id] != row:
            raise ValueError(f"同一 流水ID 出现冲突内容: {flow_id}")
        by_id[flow_id] = row
    return list(by_id.values())


def sum_amount(rows: Iterable[dict[str, str]]) -> str:
    total = sum((decimal_value(row.get("订单金额")) for row in rows), Decimal("0"))
    return format(total, "f")


def load_json_optional(path_value: str | None, default: Any) -> Any:
    if not path_value:
        return default
    return json.loads(Path(path_value).read_text(encoding="utf-8"))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--system-file", required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--order-export", action="append", required=True, help="Repeat, pass a glob, or pass a directory")
    parser.add_argument("--flow-export", action="append", required=True, help="Repeat, pass a glob, or pass a directory")
    parser.add_argument("--start-date", required=True)
    parser.add_argument("--end-date", required=True)
    parser.add_argument("--flow-start-date", help="Actual inclusive flow export start; defaults to order start")
    parser.add_argument("--flow-end-date", help="Actual inclusive flow export end; defaults to min(order end + 3 days, business today)")
    parser.add_argument("--currency", required=True)
    parser.add_argument("--business-time-zone", required=True, help="Value returned by mc_list_supported_currencies")
    parser.add_argument("--user-time-zone", default="Asia/Shanghai")
    parser.add_argument("--as-of-utc")
    parser.add_argument("--funding-domain", default="UNPAYOUT_AMT")
    parser.add_argument("--account-type", default="ORDER_AMT")
    parser.add_argument("--order-export-meta")
    parser.add_argument("--flow-export-meta")
    parser.add_argument("--output", required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    start = date.fromisoformat(args.start_date)
    end = date.fromisoformat(args.end_date)
    if start > end:
        raise ValueError("开始日期不能晚于结束日期")
    system_path = Path(args.system_file).resolve()
    order_paths = expand_paths(args.order_export)
    flow_paths = expand_paths(args.flow_export)
    system_rows = read_table(system_path, args.sheet)
    order_rows = load_many(order_paths)
    flow_rows_all = load_many(flow_paths)
    system_headers = require_columns(system_rows, SYSTEM_REQUIRED, "系统订单文件")
    order_headers = require_columns(order_rows, TDAYPAY_REQUIRED, "TDayPay 订单导出")
    flow_headers = require_columns(flow_rows_all, FLOW_REQUIRED, "TDayPay 记账流水导出")

    start_key, end_key = start.strftime("%Y%m%d"), end.strftime("%Y%m%d")
    currency = args.currency.upper()
    business_zone = ZoneInfo(args.business_time_zone)
    user_zone = ZoneInfo(args.user_time_zone)
    as_of = datetime.fromisoformat(args.as_of_utc.replace("Z", "+00:00")) if args.as_of_utc else datetime.now(timezone.utc)
    if as_of.tzinfo is None:
        as_of = as_of.replace(tzinfo=timezone.utc)
    as_of = as_of.astimezone(timezone.utc)
    business_today = as_of.astimezone(business_zone).date()
    if end > business_today:
        raise ValueError(f"结束日期 {end} 晚于 {args.business_time_zone} 当前日期 {business_today}")
    target_flow_end = end + timedelta(days=3)
    flow_start = date.fromisoformat(args.flow_start_date) if args.flow_start_date else start
    flow_end = date.fromisoformat(args.flow_end_date) if args.flow_end_date else min(target_flow_end, business_today)
    if flow_start > start:
        raise ValueError(f"流水开始日期 {flow_start} 晚于订单开始日期 {start}")
    if flow_end < end:
        raise ValueError(f"流水结束日期 {flow_end} 早于订单结束日期 {end}")
    if flow_end > business_today:
        raise ValueError(f"流水结束日期 {flow_end} 晚于 {args.business_time_zone} 当前日期 {business_today}")

    ours, excluded_ours = filter_orders(system_rows, currency, start_key, end_key, "系统订单")
    theirs, excluded_theirs = filter_orders(order_rows, currency, start_key, end_key, "TDayPay订单")
    if not ours:
        raise ValueError("系统订单文件筛选后没有范围内的 SUCCESS 订单")
    ensure_unique(ours, "平台订单号", "系统成功订单")
    ensure_unique(theirs, "平台订单号", "TDayPay成功订单")

    flow_domain = args.funding_domain.upper()
    account_type = args.account_type.upper()
    selected_flows = [
        row for row in flow_rows_all
        if clean(row.get("资金域")).upper() == flow_domain
        and clean(row.get("账户类型")).upper() == account_type
    ]
    selected_flows = deduplicate_flows(selected_flows)

    our_map = {row["平台订单号"]: row for row in ours}
    their_map = {row["平台订单号"]: row for row in theirs}
    matched_ids = sorted(our_map.keys() & their_map.keys())
    ours_only = [our_map[key] for key in sorted(our_map.keys() - their_map.keys())]
    theirs_only = [their_map[key] for key in sorted(their_map.keys() - our_map.keys())]
    matches: list[dict[str, str]] = []
    field_discrepancies: list[dict[str, str]] = []
    for platform_id in matched_ids:
        ours_row, theirs_row = our_map[platform_id], their_map[platform_id]
        checks = {
            "商户订单号": clean(ours_row.get("商户订单号")) == clean(theirs_row.get("商户订单号")),
            "订单金额": decimal_value(ours_row.get("订单金额")) == decimal_value(theirs_row.get("订单金额")),
            "币种": clean(ours_row.get("币种")) == clean(theirs_row.get("币种")),
        }
        if "支付类型" in ours_row and clean(ours_row.get("支付类型")):
            checks["支付类型"] = normalize_order_type(ours_row.get("支付类型")) == normalize_order_type(theirs_row.get("支付类型"))
        failed = [name for name, passed in checks.items() if not passed]
        match_row = {
            "平台订单号": platform_id,
            "商户订单号": ours_row.get("商户订单号", ""),
            "我方金额": normalize_money(ours_row.get("订单金额")),
            "TDayPay金额": normalize_money(theirs_row.get("订单金额")),
            "金额差": format(decimal_value(ours_row.get("订单金额")) - decimal_value(theirs_row.get("订单金额")), "f"),
            "币种": ours_row.get("币种", ""),
            "支付类型": normalize_order_type(ours_row.get("支付类型") or theirs_row.get("支付类型")),
            "字段核查结果": "全部一致" if not failed else "不一致：" + "、".join(failed),
        }
        matches.append(match_row)
        if failed:
            field_discrepancies.append(match_row)

    settlement_status_counts = Counter(clean(row.get("结算状态")).upper() for row in theirs)
    unsettled = [row for row in theirs if clean(row.get("结算状态")).upper() == "UNCHECK"]
    settlement_unknown = [
        row for row in theirs
        if clean(row.get("结算状态")).upper() not in {"CHECKED", "UNCHECK"}
    ]
    flows_by_order: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in selected_flows:
        flows_by_order[clean(row.get("订单号"))].append(row)
    matched_flows: list[dict[str, str]] = []
    out_of_window_flows: list[dict[str, str]] = []
    missing_flows: list[dict[str, str]] = []
    pending_flow_observation: list[dict[str, str]] = []
    for order in theirs:
        order_date = datetime.strptime(normalize_date(order.get("创建日期")), "%Y%m%d").date()
        window_end = order_date + timedelta(days=3)
        eligible: list[dict[str, str]] = []
        for flow in flows_by_order.get(order["平台订单号"], []):
            flow_date = datetime.strptime(normalize_date(flow.get("创建时间")), "%Y%m%d").date()
            delay_days = (flow_date - order_date).days
            enriched = dict(flow)
            enriched["关联成功订单创建日期"] = order.get("创建日期", "")
            enriched["关联商户订单号"] = order.get("商户订单号", "")
            enriched["流水距订单创建天数"] = str(delay_days)
            if 0 <= delay_days <= 3:
                enriched["订单流水窗口结果"] = "窗口内"
                eligible.append(enriched)
                matched_flows.append(enriched)
            else:
                enriched["订单流水窗口结果"] = "窗口外"
                out_of_window_flows.append(enriched)
        if eligible:
            continue
        if flow_end >= window_end:
            missing = dict(order)
            missing["流水核查结果"] = f"订单创建日起至后 3 天未查到 {flow_domain} + {account_type} 流水"
            missing_flows.append(missing)
        else:
            pending = dict(order)
            pending["流水核查结果"] = f"观察期未结束；需查询至 {window_end.isoformat()}，当前流水仅到 {flow_end.isoformat()}"
            pending_flow_observation.append(pending)

    summary = {
        "ours_count": len(ours), "ours_amount": sum_amount(ours),
        "tdaypay_success_count": len(theirs), "tdaypay_success_amount": sum_amount(theirs),
        "matched_count": len(matched_ids),
        "ours_only_count": len(ours_only), "ours_only_amount": sum_amount(ours_only),
        "tdaypay_only_count": len(theirs_only), "tdaypay_only_amount": sum_amount(theirs_only),
        "field_discrepancy_count": len(field_discrepancies),
        "amount_delta": format(Decimal(sum_amount(ours)) - Decimal(sum_amount(theirs)), "f"),
        "settled_count": sum(1 for row in theirs if clean(row.get("结算状态")).upper() == "CHECKED"),
        "unsettled_count": len(unsettled), "unsettled_amount": sum_amount(unsettled),
        "settlement_unknown_count": len(settlement_unknown),
        "settlement_unknown_amount": sum_amount(settlement_unknown),
        "settlement_status_counts": dict(sorted(settlement_status_counts.items())),
        "flow_raw_count": len(selected_flows), "matched_flow_count": len(matched_flows),
        "out_of_window_flow_count": len(out_of_window_flows),
        "covered_order_count": len(theirs) - len(missing_flows) - len(pending_flow_observation),
        "missing_flow_count": len(missing_flows), "missing_flow_amount": sum_amount(missing_flows),
        "pending_flow_observation_count": len(pending_flow_observation),
        "pending_flow_observation_amount": sum_amount(pending_flow_observation),
        "payment_count": sum(1 for row in theirs if normalize_order_type(row.get("支付类型")) == "PAYMENT"),
        "payout_count": sum(1 for row in theirs if normalize_order_type(row.get("支付类型")) == "PAYOUT"),
        "dates": dict(sorted(Counter(row["创建日期"] for row in theirs).items())),
    }
    meta = {
        "source_file": str(system_path),
        "currency": currency,
        "business_time_zone": args.business_time_zone,
        "user_time_zone": args.user_time_zone,
        "period_start": start.isoformat(),
        "period_end": end.isoformat(),
        "period_complete": end < business_today,
        "flow_period_start": flow_start.isoformat(),
        "flow_period_end": flow_end.isoformat(),
        "flow_target_end": target_flow_end.isoformat(),
        "flow_window_complete": flow_end >= target_flow_end,
        "as_of_utc": as_of.isoformat().replace("+00:00", "Z"),
        "as_of_business": as_of.astimezone(business_zone).isoformat(),
        "as_of_user": as_of.astimezone(user_zone).isoformat(),
        "flow_funding_domain": flow_domain,
        "flow_account_type": account_type,
        "order_export_files": [str(path) for path in order_paths],
        "flow_export_files": [str(path) for path in flow_paths],
        "order_exports": load_json_optional(args.order_export_meta, []),
        "flow_exports": load_json_optional(args.flow_export_meta, []),
    }
    output = {
        "meta": meta,
        "summary": summary,
        "system_headers": system_headers,
        "order_headers": order_headers,
        "flow_headers": flow_headers + ["关联成功订单创建日期", "关联商户订单号", "流水距订单创建天数", "订单流水窗口结果"],
        "ours": ours,
        "success": theirs,
        "matches": matches,
        "field_discrepancies": field_discrepancies,
        "only_ours": ours_only,
        "only_theirs": theirs_only,
        "unsettled": unsettled,
        "settlement_unknown": settlement_unknown,
        "raw_flows": selected_flows,
        "matched_flows": matched_flows,
        "out_of_window_flows": out_of_window_flows,
        "missing": missing_flows,
        "pending_flow_observation": pending_flow_observation,
        "excluded_system_rows": excluded_ours,
        "excluded_tdaypay_rows": excluded_theirs,
    }
    output_path = Path(args.output).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output_path), "summary": summary}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(2)
