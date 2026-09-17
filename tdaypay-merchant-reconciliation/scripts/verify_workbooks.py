"""Verify the six text-only workbooks produced by build_workbooks.js."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")


def normalized(value: Any) -> str:
    text = "" if value is None else str(value)
    return text[1:] if text.startswith("'") else text


def expected_names(data: dict[str, Any]) -> list[str]:
    meta = data["meta"]
    currency = meta["currency"]
    start = meta["period_start"].replace("-", "")
    end = meta["period_end"].replace("-", "")
    domain = meta["flow_funding_domain"]
    account = meta["flow_account_type"]
    return [
        f"{currency}_对账总结_{start}-{end}.xlsx",
        f"{currency}_差异订单_{start}-{end}.xlsx",
        f"{currency}_未结算订单_{start}-{end}.xlsx",
        f"TDayPay_{currency}_成功订单_{start}-{end}.xlsx",
        f"{currency}_{domain}_{account}_未查到流水订单_{start}-{end}.xlsx",
        f"TDayPay_{currency}_{domain}_{account}_记账流水_{start}-{end}.xlsx",
    ]


def verify_all_cells_are_text(path: Path) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    if getattr(workbook, "_external_links", []):
        raise AssertionError(f"{path.name} 包含外部链接")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if cell.data_type != "s":
                    raise AssertionError(f"{path.name}#{sheet.title}!{cell.coordinate} 不是文本单元格: {cell.data_type}")
                if any(token in str(cell.value) for token in ERROR_TOKENS):
                    raise AssertionError(f"{path.name}#{sheet.title}!{cell.coordinate} 含错误值: {cell.value}")


def verify_detail(path: Path, sheet_name: str, expected_rows: list[dict[str, Any]], expected_headers: list[str]) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheet = workbook[sheet_name]
    actual_headers = [normalized(cell.value) for cell in sheet[9]][:len(expected_headers)]
    if actual_headers != expected_headers:
        raise AssertionError(f"{path.name}#{sheet_name} 表头不一致: {actual_headers} != {expected_headers}")
    if not expected_rows:
        if normalized(sheet["A10"].value) not in {
            "无符合条件的记录", "无差异订单", "无缺少流水的成功订单",
            "无观察期未结束的订单", "无订单号匹配但超窗的流水",
            "无结算状态缺失或未知的成功订单",
        }:
            raise AssertionError(f"{path.name}#{sheet_name} 空结果标记异常: {sheet['A10'].value}")
        return
    for row_number, expected in enumerate(expected_rows, 10):
        for column_number, key in enumerate(expected_headers, 1):
            actual = normalized(sheet.cell(row_number, column_number).value)
            wanted = normalized(expected.get(key, ""))
            if actual != wanted:
                raise AssertionError(
                    f"{path.name}#{sheet_name}!{sheet.cell(row_number, column_number).coordinate} "
                    f"值不一致: {actual!r} != {wanted!r} ({key})"
                )
    extra_row = 10 + len(expected_rows)
    if any(cell.value is not None for cell in sheet[extra_row]):
        raise AssertionError(f"{path.name}#{sheet_name} 存在未预期的额外数据行 {extra_row}")


def verify_summary(path: Path, summary: dict[str, Any]) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheet = workbook["对账概览"]
    values = {normalized(sheet.cell(row, 1).value): normalized(sheet.cell(row, 2).value) for row in range(6, sheet.max_row + 1)}
    checks = {
        "我方成功订单": str(summary["ours_count"]),
        "TDayPay 成功订单": str(summary["tdaypay_success_count"]),
        "双方匹配订单": str(summary["matched_count"]),
        "我方有、TDayPay 无": str(summary["ours_only_count"]),
        "TDayPay 有、我方无": str(summary["tdaypay_only_count"]),
        "字段差异订单": str(summary["field_discrepancy_count"]),
        "TDayPay 未结算笔数": str(summary["unsettled_count"]),
        "结算状态缺失或未知笔数": str(summary.get("settlement_unknown_count", 0)),
        "未查到指定流水的成功订单": str(summary["missing_flow_count"]),
        "流水观察期未结束订单": str(summary.get("pending_flow_observation_count", 0)),
        "指定口径流水条数": str(summary["flow_raw_count"]),
        "订单号匹配但超窗流水": str(summary.get("out_of_window_flow_count", 0)),
    }
    for label, wanted in checks.items():
        if values.get(label) != wanted:
            raise AssertionError(f"{path.name} 汇总项 {label} 不一致: {values.get(label)!r} != {wanted!r}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--report", help="Optional path for a machine-readable verification report")
    args = parser.parse_args()
    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    output_dir = Path(args.output_dir).resolve()
    names = expected_names(data)
    actual_names = sorted(path.name for path in output_dir.glob("*.xlsx"))
    if actual_names != sorted(names):
        raise AssertionError(f"工作簿集合不一致: {actual_names} != {sorted(names)}")
    for name in names:
        verify_all_cells_are_text(output_dir / name)

    currency = data["meta"]["currency"]
    start = data["meta"]["period_start"].replace("-", "")
    end = data["meta"]["period_end"].replace("-", "")
    domain = data["meta"]["flow_funding_domain"]
    account = data["meta"]["flow_account_type"]
    summary_path = output_dir / f"{currency}_对账总结_{start}-{end}.xlsx"
    differences_path = output_dir / f"{currency}_差异订单_{start}-{end}.xlsx"
    unsettled_path = output_dir / f"{currency}_未结算订单_{start}-{end}.xlsx"
    orders_path = output_dir / f"TDayPay_{currency}_成功订单_{start}-{end}.xlsx"
    missing_path = output_dir / f"{currency}_{domain}_{account}_未查到流水订单_{start}-{end}.xlsx"
    flows_path = output_dir / f"TDayPay_{currency}_{domain}_{account}_记账流水_{start}-{end}.xlsx"

    verify_summary(summary_path, data["summary"])
    verify_detail(summary_path, "对账明细", data["matches"], ["平台订单号", "商户订单号", "我方金额", "TDayPay金额", "金额差", "币种", "支付类型", "字段核查结果"])
    verify_detail(differences_path, "我方有TDayPay无", data["only_ours"], data["system_headers"])
    verify_detail(differences_path, "TDayPay有我方无", data["only_theirs"], data["order_headers"])
    verify_detail(unsettled_path, "明确未结算", data["unsettled"], data["order_headers"])
    verify_detail(unsettled_path, "结算状态待核实", data.get("settlement_unknown", []), data["order_headers"])
    verify_detail(orders_path, "成功订单", data["success"], data["order_headers"])
    verify_detail(missing_path, "未查到流水订单", data["missing"], [*data["order_headers"], "流水核查结果"])
    verify_detail(missing_path, "观察期未结束", data.get("pending_flow_observation", []), [*data["order_headers"], "流水核查结果"])
    verify_detail(flows_path, "成功订单关联流水", data["matched_flows"], data["flow_headers"])
    verify_detail(flows_path, "订单号匹配但超窗", data.get("out_of_window_flows", []), data["flow_headers"])
    raw_headers = [header for header in data["flow_headers"] if header not in {"关联成功订单创建日期", "关联商户订单号", "流水距订单创建天数", "订单流水窗口结果"}]
    verify_detail(flows_path, "全部查询流水", data["raw_flows"], raw_headers)

    for row in data["raw_flows"]:
        if row.get("资金域", "").upper() != domain or row.get("账户类型", "").upper() != account:
            raise AssertionError(f"流水超出口径: {row.get('流水ID')}")
    success_ids = {row["平台订单号"] for row in data["success"]}
    covered_ids = {row["订单号"] for row in data["matched_flows"]}
    missing_ids = {row["平台订单号"] for row in data["missing"]}
    pending_ids = {row["平台订单号"] for row in data.get("pending_flow_observation", [])}
    if success_ids != covered_ids | missing_ids | pending_ids or covered_ids & missing_ids or covered_ids & pending_ids or missing_ids & pending_ids:
        raise AssertionError("流水覆盖订单集合与成功订单集合不一致")
    for row in data["matched_flows"]:
        if not 0 <= int(row["流水距订单创建天数"]) <= 3 or row["订单流水窗口结果"] != "窗口内":
            raise AssertionError(f"窗口内流水分类错误: {row.get('流水ID')}")
    for row in data.get("out_of_window_flows", []):
        if 0 <= int(row["流水距订单创建天数"]) <= 3 or row["订单流水窗口结果"] != "窗口外":
            raise AssertionError(f"超窗流水分类错误: {row.get('流水ID')}")

    report = {
        "Verification": "PASS",
        "files": names,
        "all_nonblank_cells_text": True,
        "formula_errors": 0,
        "external_links": 0,
        "system_success_orders": data["summary"]["ours_count"],
        "tdaypay_success_orders": data["summary"]["tdaypay_success_count"],
        "unsettled_orders": data["summary"]["unsettled_count"],
        "settlement_unknown_orders": data["summary"].get("settlement_unknown_count", 0),
        "missing_flow_orders": data["summary"]["missing_flow_count"],
        "pending_flow_observation_orders": data["summary"].get("pending_flow_observation_count", 0),
        "out_of_window_flows": data["summary"].get("out_of_window_flow_count", 0),
    }
    if args.report:
        report_path = Path(args.report).resolve()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("Verification: PASS")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"Verification: FAIL\n{error}", file=sys.stderr)
        raise SystemExit(2)
